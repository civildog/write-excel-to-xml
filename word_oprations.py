import win32com.client
import os
import time
import sys
import docx


path_1 = r'c:\users\admin\desktop\31  就地浇筑梁、板现场质量检验表(检表8.7.1)(SG)-有试验检测评定表关系.doc'
path = r'c:\users\admin\desktop\python学习笔记.docx'
path_2 = r'c:\users\admin\desktop\临时文件'
tem_txt = r'c:\users\admin\desktop\table.txt'


def doc_to_docx(path,to_path):
    """将doc格式转化为docx格式"""
    try:
        myword = win32com.client.DispatchEx('Word.Application')
        myword.Visible= 0
        myword.DisplayAlerts = 0
        doc = myword.Documents.Open(path)

        name = os.path.basename(path).rstrip('.doc')+'.docx'
        doc.SaveAs(os.path.join(to_path, name), FileFormat=16) # FileFormat 16 保存成docx文件，17保存成pdf文件
    finally:
        doc.Close()
        myword.Quit()


# f = docx.Document(path)
# t = f.tables[0]
# for row in range(len(t.rows)):
#     for col in range(len(t.columns)):
#         if ('规定值或' and '允许偏差') in t.cell(row, col).text:
#             print(col)
#
# for row in range(len(t.rows)):
#     for col in range(len(t.columns)):
#         if '桥梁记录表' in t.cell(row, col).text:
#             print('row:{0},col:{1}'.format(row,col))
#             print(t.cell(row, 3).text)
# myword = win32com.client.DispathEx('Word.Application')
# new = myword.Documents.Add()
# new.SaveAs(to_path)
# range = doc.Content
# range = doc.Range(doc.Content.Start, doc.Content.End)
# range.text = 'haha'
# selection = range.Table()
# print(range)
# table = doc.Tables
# x = doc.Range().Tables(1).range.start
# f = doc.Paragraphs(doc.Range(0, x).Paragraphs.Count).Range()
# f = doc.Paragraphs(10).Range.text
# f = doc.Tables(1).Columns.Count
# t = doc.Tables(1)
# selection = doc.select(t.Cell(10,4).Range.start, t.Cell(10,4).Range.end).Text
# print(str(selection))
# row_num = t.Rows.Count
# col_num = t.Columns.Count
# for row in range(1, row_num+1):
#     for col in range(1, col_num+1):
#         if "测表3" in str(doc.Tables(1).Cell(row, col).Range.Text):
#             print("True")
# f = t.Cell(1,1)
# f = t.Cell(1,1).Range.Select()
# print(f)
# t = open(tem_txt, 'w')
# t.write(f)
# f = doc.Content.Find.Execute(FindText='测表3')
# t.Range.MoveStart(t.Cell(10,4),1)
import re
if __name__ == "__main__":
    line = '路面记录表16（索引表中所有桩号的“实测值”所有列“数据”）'
    # line = 'asdvd'
    a = re.match('.*记录表.*', line)
    # a = re.search('[a-z]+',line)
    start = re.search('“', line).span()
    start2 = re.search('”', line).span()
    start3 = re.search('aaaa', line)
    t= '高强度（Nm)'
    print(re.sub('[(（][a-zA-Z]+.?[a-zA-Z]+[)）]','',t))
    # s = '架设拱圈前，台后沉降完成量（mm）'
    # print(re.sub('（[a-z]*）','',s))
    # n = '%TCR23'
    # if re.match('%\w{3,4}\d{1,2}',n):
    #     print(re.match('%[a-zA-Z]{3,4}\d{1,2}',n).group())
    # header = '架设拱圈前，台后沉降完成量（mm）'
    # print(type(re.sub('（[a-z]*）', '', header)))
    # t = '见测表3（索引“偏差”一列的所有数值）'
    # print(t.split('“')[1].split('”')[0])
    # content = '换购诋毁和'
    # print(content.split('、'))
    # o = "见路面记录表03（索引表中“实测厚度、偏差”所在列所有数据）(将“实测厚度”放在实测值里面，将“偏差”放在偏差值里面)"
    # w = o.split('“')[1].split('”')[0].split('、')

