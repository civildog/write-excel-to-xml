# encoding:utf-8
import time
import os
import re
import win32com.client
import docx
import xml.dom.minidom
import json

# path = r'C:\Users\admin\Desktop\datarelation\result\2019年四川省公路工程施工及监理统一用表汇编\4 桥梁工程(第四册)\1 桥梁工程\1 施工单位\2 检表\桥梁检表\29  拱桥组合桥台现场质量检验表(检表8.6.3)(SG).xlsx'
path = r'C:\Users\admin\Desktop\datarelation\result\2019年四川省公路工程施工及监理统一用表汇编\3 路面工程(第三册)\1 路面工程\1 施工单位\1 评定表\路面评定表\评定表7.2.2  水泥混凝土面层评定(SG).xlsx'
temp_path = r'C:\Users\admin\Desktop\datarelation\tempfile'
log_path = r'C:\Users\admin\Desktop\datarelation\tempfile\log'
xml_path = r'C:\Users\admin\Desktop\datarelation\set.xml'


def find_adr(path, content):
    """查找存在内容content的地址，并返回[row, col]"""
    adr_list = []
    if os.path.exists(path):
        if os.path.isfile(path):
            try:
                excel = win32com.client.DispatchEx("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = 0
                work = excel.Workbooks.Open(path)
                sheet = work.Worksheets(1)
                cel = sheet.UsedRange.Find(content)
                if cel:
                    adr = cel.Address
                    col_row = adr.split("$")
                    row = col_row[2]
                    adr_list.append(int(row))
                    col = ord(col_row[1]) - 64
                    adr_list.append(col)
                    return adr_list
                else:
                    return None
            except:
                print("无法找到指定内容！")
            finally:
                work.Close()
                excel.Quit()
        else:
            print("不是文件！")
    else:
        print("目录不存在！")



def doc_to_docx(path,to_path):
    """将doc格式转化为docx格式"""
    try:
        myword = win32com.client.DispatchEx('Word.Application')
        myword.Visible= 0
        myword.DisplayAlerts = 0
        doc = myword.Documents.Open(path)

        name = os.path.basename(path).rstrip('.doc')+'.docx'
        doc.SaveAs(os.path.join(to_path, name), FileFormat=16) # FileFormat 16 保存成docx文件，17保存成pdf文件
    finally:
        doc.Close()
        myword.Quit()


def get_info(path, tempath,log_path):
    """ 获取数据关系信息 """
    new_path = os.path.join(tempath, os.path.splitext(path)[0].split('\\')[-1] + '.docx')
    result = {'测表3': None, '测表6': None, '测表11': None, '测表16': None, '关系表': None}
    ce3, ce6, ce11, ce16 = {'rows': [], 'content': []}, {'rows': [], 'content': []}, \
                           {'rows': [], 'content': []}, {'rows': [], 'content': []}
    target = {'rows': [], 'content': []}
    word = docx.Document(new_path)
    t = word.tables[0]
    header_col = None
    for row in range(len(t.rows)):
        content3 = {'row': None, 'cols': [], 'header': None, 'info': ''}
        content6 = {'row': None, 'cols': [], 'header': None, 'info': ''}
        content11 = {'row': None, 'cols': [], 'header': None, 'info': ''}
        content16 = {'row': None, 'cols': [], 'header': None, 'info': ''}
        content_target = {'row': None, 'cols': [], 'header': None, 'targetfile': None, 'info': ''}
        for col in range(len(t.columns)):
            if (re.match('规定值.{4,6}',re.sub('\s*','',t.cell(row, col).text))) and not header_col:
                header_col = col
            if '测表3' in t.cell(row, col).text:
                if row not in ce3['rows']:
                    ce3['rows'].append(row)
                    content3['row'] = row
                    content3['info'] = re.sub('\s*','',t.cell(row, col).text)
                if col not in content3['cols']:
                    content3['cols'].append(col)
                break
            elif '测表6' in t.cell(row, col).text:
                if row not in ce6['rows']:
                    ce6['rows'].append(row)
                    content6['row'] = row
                    content6['info'] = re.sub('\s*','',t.cell(row, col).text)
                if col not in content6['cols']:
                    content6['cols'].append(col)
                break
            elif '测表11' in t.cell(row, col).text:
                if row not in ce11['rows']:
                    ce11['rows'].append(row)
                    content11['row'] = row
                    content11['info'] = re.sub('\s*','',t.cell(row, col).text)
                if col not in content11['cols']:
                    content11['cols'].append(col)
                break
            elif '测表16' in t.cell(row, col).text:
                if row not in ce16['rows']:
                    ce16['rows'].append(row)
                    content16['row'] = row
                    content16['info'] = re.sub('\s*','',t.cell(row, col).text)
                if col not in content16['cols']:
                    content16['cols'].append(col)
                break
            elif re.match('.*记录表.*', t.cell(row, col).text):
                if row not in target['rows']:
                    target['rows'].append(row)
                    target_name = t.cell(row, col).text.split('（')[0].lstrip('见')
                    content_target['row'] = row
                    content_target['targetfile'] = target_name
                    content_target['info'] = re.sub('\s*','',t.cell(row, col).text)
                if col not in content_target['cols']:
                    content_target['cols'].append(col)
                break
        if content3['row']:
            content3['header'] = re.sub('\s*','',t.cell(row, header_col - 1).text)
            ce3['content'].append(content3)
        if content6['row']:
            content6['header'] = re.sub('\s*','',t.cell(row, header_col - 1).text)
            ce6['content'].append(content6)
        if content11['row']:
            content11['header'] = re.sub('\s*','',t.cell(row, header_col - 1).text)
            ce11['content'].append(content11)
        if content16['row']:
            content16['header'] = re.sub('\s*','',t.cell(row, header_col - 1).text)
            ce16['content'].append(content16)
        if content_target['row']:
            content_target['header'] = re.sub('\s*','',t.cell(row, header_col - 1).text)
            target['content'].append(content_target)

    f = open(os.path.join(log_path, time.strftime('%Y-%m-%d', time.localtime()) + '.txt'), 'a',encoding='utf-8')
    f.write('\n' + time.strftime('%H:%M:%S', time.localtime()) + ":" + '\n\t' + os.path.basename(path) + '\n\t\t' +
            '测表3:{}'.format(str(ce3)) + '\n\t\t' + '测表6:{}'.format(str(ce6)) + '\n\t\t' + '测表11:{}'.format(str(ce11)) +
            '\n\t\t' + '测表16:{}'.format(str(ce16)) + '\n\t\t' + '记录表数据来源:{}'.format(str(target)))
    f.close()

    result['测表3'], result['测表6'], result['测表11'], result['测表16'], result['关系表'] = ce3, ce6, ce11, ce16, target
    return result


def read_path(xml_path, name):
    """读取set文件路径，找到每一张excel表对应的xml"""
    result = {}
    dom = xml.dom.minidom.parse(xml_path)
    root = dom.documentElement
    tabs = root.getElementsByTagName('Tab')
    for tab in tabs:
        excel_name = tab.getAttribute('ExcelName')
        if name in excel_name:
            id = tab.getAttribute('TabId')
            xml_name = tab.getAttribute('XmlName')
            result['Id'] = id
            result['ExcelName'] = excel_name
            result['XmlName'] = xml_name
        # print(id+'\n'+excel_name+'\n'+xml_name)
    return result


def inspect_info(path, tempath, log_path):
    """解析检表与其他表单间的关系，并且将其按照规定的格式写入到对应的info当中"""
    # 转化路径
    pre_path = os.path.splitext(path)[0]+'.doc'
    info_path = re.sub('result', 'relation', pre_path)
    # 先将doc转化成docx
    sub_name = os.path.basename(info_path).rstrip('.doc')+'.docx'
    files = os.listdir(tempath)
    if sub_name not in files:
        doc_to_docx(info_path, tempath)
    else:
        pass
    # 根据表单查询数据描述的文档，并提取信息
    info = get_info(info_path, tempath, log_path)
    log = open(os.path.join(log_path, time.strftime('%Y-%m-%d', time.localtime()) + 'log.txt'), 'a',encoding='utf-8')
    log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'解析：{}'.format(os.path.basename(path))+'\n')
    # 根据数据关系信息查找对应表单的key
    # 查询本表被填充的key
    try:
        myexcel = win32com.client.DispatchEx('Excel.Application')
        myexcel.Visible = 0
        myexcel.DisplayAlerts = 0
        work = myexcel.Workbooks.Open(path)
        sheet = work.Worksheets(1)
        rows = sheet.UsedRange.Rows.Count
        cols = sheet.UsedRange.Columns.Count
        if len(info['测表3']['rows']) > 0:
            for i in range(len(info['测表3']['content'])):
                log.write(time.strftime('%H:%M:%S',time.localtime())+':\t'+'解析：测表3 info'+'\n')
                for row in range(1, rows+1):
                    for col in range(1, cols+1):
                        # a = re.sub('（[a-z]*）', '', info['测表3']['content'][i]['header'])
                        a = info['测表3']['content'][i]['header']
                        name = sheet.Cells(row, col).Value
                        if re.sub('[（(][a-zA-Z]*[)）]','',a) == re.sub('[（(][a-zA-Z]*[)）]','',re.sub('\s*', '', str(sheet.Cells(row, col).Value))):
                            for j in range(col, cols+1):
                                if re.match('%\w{3,4}\d{1,2}',str(sheet.Cells(row, j).Value)):
                                    log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入to_key'+'\n')
                                    info['测表3']['content'][i]['to_key'] = sheet.Cells(row, j).Value[3:]
                                    break
                describe = info['测表3']['content'][i]['info'].split('“')[1].split('”')[0]
                if ('偏差' and '高程') in describe:
                    if '(SG)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入(SG)from_key(2个)'+'\n')
                        info['测表3']['content'][i]['from_key'] = 'LBZJTY0069!%TCQ09:%TCQ33,LBZJTY0069!%TCW09:%TCW33'
                    elif '(JL)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' +'写入(JL)from_key(2个)'+'\n')
                        info['测表3']['content'][i]['from_key'] = 'LBZJTY0092!%TCQ09:%TCQ33,LBZJTY0092!%TCW09:%TCW33'
                elif '偏差' in describe:
                    if '(SG)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(SG)from_key'+'\n')
                        info['测表3']['content'][i]['from_key'] = 'LBZJTY0069!%TCW09:%TCW33'
                    elif '(JL)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(JL)from_key'+'\n')
                        info['测表3']['content'][i]['from_key'] = 'LBZJTY0092!%TCW09:%TCW33'
        if len(info['测表6']['rows']) > 0:
            for i in range(len(info['测表6']['content'])):
                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'解析：测表6 info'+'\n')
                for row in range(1, rows+1):
                    for col in range(1, cols+1):
                        a = info['测表6']['content'][i]['header']
                        if re.sub('[（(][a-zA-Z]*[)）]','',a) == re.sub('[（(][a-zA-Z]*[)）]','',re.sub('\s*', '', str(sheet.Cells(row, col).Value))):
                            for j in range(col, cols+1):
                                if re.match('%\w{3,4}\d{1,2}',str(sheet.Cells(row, j).Value)):
                                    info['测表6']['content'][i]['to_key'] = sheet.Cells(row, j).Value[3:]
                                    break
                describe = info['测表6']['content'][i]['info'].split('“')[1].split('”')[0]
                if '宽度-左右侧的宽度' or '宽度 -实测值-左、右' in describe:
                    if '(SG)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(SG)from_key'+'\n')
                        info['测表6']['content'][i]['from_key'] = 'LBZJTY0072!%TCG27:%TCY28'
                    elif '(JL)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(JL)from_key'+'\n')
                        info['测表6']['content'][i]['from_key'] = 'LBZJTY0095!%TCG27:%TCY28'
                elif '纵段高程-左右侧的宽度' or '纵断高程-差值-左边、中桩、右边' in describe:
                    if '(SG)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(SG)from_key'+'\n')
                        info['测表6']['content'][i]['from_key'] = 'LBZJTY0072!%TCG16:%TCY18'
                    elif '(JL)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(JL)from_key'+'\n')
                        info['测表6']['content'][i]['from_key'] = 'LBZJTY0095!%TCG16:%TCY18'
                elif '横坡度-左右侧的宽度' or '横坡度-偏差值-左、右偏差值' in describe:
                    if '(SG)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(SG)from_key'+'\n')
                        info['测表6']['content'][i]['from_key'] = 'LBZJTY0072!%TCG23:%TCY24'
                    elif '(JL)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(JL)from_key'+'\n')
                        info['测表6']['content'][i]['from_key'] = 'LBZJTY0095!%TCG16:%TCY18'
        if len(info['测表11']['rows']) > 0:
            for i in range(len(info['测表11']['content'])):
                log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t' +'解析：测表11 info'+'\n')
                for row in range(1, rows+1):
                    for col in range(1, cols+1):
                        a = info['测表11']['content'][i]['header']
                        if re.sub('[（(][a-zA-Z]*[)）]','',a) == re.sub('[（(][a-zA-Z]*[)）]','',re.sub('\s*', '', str(sheet.Cells(row, col).Value))):
                            for j in range(col, cols+1):
                                if re.match('%\w{3,4}\d{1,2}',str(sheet.Cells(row, j).Value)):
                                    log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入to_key'+'\n')
                                    info['测表11']['content'][i]['to_key'] = sheet.Cells(row, j).Value[3:]
                                    break
                describe = info['测表11']['content'][i]['info'].split('“')[1].split('”')[0]
                if '偏位' in describe:
                    if '(SG)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(SG)from_key'+'\n')
                        info['测表11']['content'][i]['from_key'] = 'LBZJTY0078!%TCY13:%TCY28'
                    elif '(JL)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(JL)from_key'+'\n')
                        info['测表11']['content'][i]['from_key'] = 'LBZJTY0101!%TCY13:%TCY28'
        if len(info['测表16']['rows']) > 0:
            for i in range(len(info['测表16']['content'])):
                log.write(time.strftime('%H:%M:%S', time.localtime()) + '解析：测表16 info'+'\n')
                for row in range(1, rows+1):
                    for col in range(1, cols+1):
                        a = info['测表16']['content'][i]['header']
                        if re.sub('[（(][a-zA-Z]*[)）]','',a) == re.sub('[（(][a-zA-Z]*[)）]','',re.sub('\s*', '', str(sheet.Cells(row, col).Value))):
                            for j in range(col, cols+1):
                                if re.match('%\w{3,4}\d{1,2}',str(sheet.Cells(row, j).Value)):
                                    log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入to_key'+'\n')
                                    info['测表16']['content'][i]['to_key'] = sheet.Cells(row, j).Value[3:]
                                    break
                describe = info['测表16']['content'][i]['info'].split('“')[1].split('”')[0]
                if '坡度' in describe:
                    if '(SG)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(SG)from_key'+'\n')
                        info['测表16']['content'][i]['from_key'] = 'LBZJTY0083!%TCH11:%TCH33,LBZJTY0083!%TCW11:%TCW33'
                    elif '(JL)' in os.path.basename(path):
                        log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入(JL)from_key'+'\n')
                        info['测表16']['content'][i]['from_key'] = 'LBZJTY0106!%TCH11:%TCH33,LBZJTY0106!%TCW11:%TCW33'
        if len(info['关系表']['rows']) > 0:
            for i in range(len(info['关系表']['content'])):
                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'解析：{}'.format(info['关系表']['content'][i]['targetfile'])+'\n')
                try:
                    for row in range(1, rows+1):
                        for col in range(1, cols+1):
                            a = info['关系表']['content'][i]['header']
                            if re.sub('[（(][a-zA-Z]*.?[A-Za-z]*[)）]','',a) == re.sub('[（(][a-zA-Z]*.?[A-Za-z]*[)）]','',re.sub('\s*', '', str(sheet.Cells(row, col).Value))):
                                for j in range(col, cols + 1):
                                    if re.match('%\w{3,4}\d{1,2}', str(sheet.Cells(row, j).Value)):
                                        log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入to_key'+'\n')
                                        info['关系表']['content'][i]['to_key'] = sheet.Cells(row, j).Value[3:]
                                        break
                    if '桥梁' in info['关系表']['content'][i]['targetfile']:
                        num = info['关系表']['content'][i]['targetfile'].lstrip('桥梁记录表')
                        root_path = os.path.dirname(os.path.dirname(os.path.dirname(path)))
                        pre_path = os.path.join(root_path, r'4 记录表\桥梁记录表')
                        file_list = os.listdir(pre_path)
                        target_path = ''
                        for f in file_list:
                            if re.match('{}.*'.format(num),f):
                                target_path = os.path.join(pre_path, f)
                                break
                        excel = win32com.client.DispatchEx('Excel.Application')
                        excel.Visible = 0
                        excel.DisplayAlerts = 0
                        work_2 = excel.Workbooks.Open(target_path)
                        sheet_2 = work_2.Worksheets(1)
                        rows_2 = sheet_2.UsedRange.Rows.Count
                        cols_2 = sheet_2.UsedRange.Columns.Count
                        set_tab = read_path(xml_path, os.path.basename(target_path))
                        keys_small = []
                        keys_large = []
                        keys = []
                        row_list = []
                        col_list = []
                        start_row = None
                        end_row = None
                        adr = [] # find_adr(target_path, info['关系表']['content'][i]['header'])
                        for row in range(6, rows_2+1):
                            for col in range(1, cols+1):
                                if re.sub('[a-zA-Z]{0,2}','',re.sub('[(（][a-zA-Z]*.?[A-Za-z]*[)）]','',info['关系表']['content'][i]['header'])) in re.sub('[a-zA-Z]{0,2}','',re.sub('[(（][a-zA-Z]*.?[A-Za-z]*[)）]','',re.sub('\s*','',str(sheet_2.Cells(row, col).Value)))):
                                    adr.append(row)
                                    adr.append(col)
                                    break
                            if adr !=[]:
                                break
                        print(info['关系表']['content'][i]['header'])
                        print(adr)
                        for r in range(adr[0]+1, rows_2+1):
                            if sheet_2.Cells(r, adr[1]).Value is not None:
                                end_row = r-1
                                break
                        for c in range(adr[1]+1, cols_2+1):
                            if re.match('%[a-zA-Z]{3,4}\d*', str(sheet_2.Cells(adr[0], c).Value)):
                                keys.append(sheet_2.Cells(adr[0], c).Value)
                                col_list.append(c)
                        '''       
                        for row in range(1, rows_2+1):
                            for col in range(1, cols_2+1):
                                if sheet_2.Cells(row, col).Value == info['关系表']['content'][i]['header']:
                                    for n in range(row+1, rows_2+1):
                                        if sheet_2.Cells(n, col).Value is not None:
                                            end_row = n-1
                                            break
                                    for j in range(col, cols_2+1):
                                        if re.match('%[a-zA-Z]{3,4}\d*', str(sheet_2.Cells(row, j).Value)):
                                            keys.append(sheet_2.Cells(row, j).Value)
                                            start_row = row
                                            col_list.append(j)
                                            # if len(sheet_2.Cells(row, j).Value[3:]) == 3:
                                            #     keys_small.append(sheet_2.Cells(row, j).Value)
                                            # elif len(sheet_2.Cells(row, j).Value[3:]) == 4:
                                            #     keys_large.append(sheet_2.Cells(row, j).Value)
                        # key_small,key_large = '',''
                        # letters, numbers = [], []
                        # for key_s in keys_small:
                        #     letters.append(key_s[3:4])
                        #     numbers.append(key_s[5:7])
                        # keys_small = None'''
                        # print(info['关系表']['content'][i]['header'])
                        key_small = sheet_2.Cells(adr[0],min(col_list)).Value
                        key_large = sheet_2.Cells(end_row,max(col_list)).Value
                        log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入关系表from_key'+'\n')
                        info['关系表']['content'][i]['from_key'] = set_tab['Id']+'!'+key_small+':'+key_large
                    elif '路面' or '路基' in info['关系表']['content'][i]['targetfile']:
                        num = info['关系表']['content'][i]['targetfile'].lstrip('路面记录表')
                        root_path = os.path.dirname(os.path.dirname(os.path.dirname(path)))
                        pre_path = ''
                        if '路面' in info['关系表']['content'][i]['targetfile']:
                            pre_path = os.path.join(root_path, r'4 记录表\路面记录表')
                        elif '路基' in info['关系表']['content'][i]['targetfile']:
                            pre_path = os.path.join(root_path, r'4 记录表\路基土石方记录表')
                        file_list = os.listdir(pre_path)
                        target_path = ''
                        for f_2 in file_list:
                            if num in f_2:
                                target_path = os.path.join(pre_path, f_2)
                                break
                        excel = win32com.client.DispatchEx('Excel.Application')
                        excel.Visible = 0
                        excel.DisplayAlerts = 0
                        work_2 = excel.Workbooks.Open(target_path)
                        sheet_2 = work_2.Worksheets(1)
                        rows_2 = sheet_2.UsedRange.Rows.Count
                        cols_2 = sheet_2.UsedRange.Columns.Count
                        set_tab = read_path(xml_path, os.path.basename(target_path))
                        describe = info['关系表']['content'][i]['info'].split('“')[1].split('”')[0].split('、')
                        for j in describe:
                            for row in range(1, rows_2+1):
                                for col in range(1, cols_2+1):
                                    if j == re.sub('\s*','',str(sheet_2.Cells(row, col).Value)):
                                        row_list = []
                                        for r in range(row + 1, rows_2 + 1):
                                            if re.match('%[a-zA-Z]{3,4}\d*', str(sheet_2.Cells(r, col).Value)):
                                                row_list.append(r)
                                        key_small = sheet_2.Cells(min(row_list), col).Value
                                        key_large = sheet_2.Cells(max(row_list), col).Value
                                        if 'from_key' in info['关系表']['content'][i].keys():
                                            log.write(time.strftime('%H:%M:%S',
                                                                    time.localtime()) + ':\t' + '写入关系表from_key' + '\n')
                                            info['关系表']['content'][i]['from_key_2'] = set_tab[
                                                                                          'Id'] + '!' + key_small + ':' + key_large
                                        else:
                                            log.write(time.strftime('%H:%M:%S',
                                                                    time.localtime()) + ':\t' + '写入关系表from_key' + '\n')
                                            info['关系表']['content'][i]['from_key'] = set_tab['Id'] + '!' + key_small + ':' + key_large
                                    # if j in re.sub('\s*','',str(sheet_2.Cells(row, col).Value)) and ('实测值' or '实测数据') not in re.sub('\s*','',str(sheet_2.Cells(row, col).Value)):
                                    #     row_list = []
                                    #     for r in range(row+1, rows_2 + 1):
                                    #         if re.match('%[a-zA-Z]{3,4}\d*', str(sheet_2.Cells(r, col).Value)):
                                    #             row_list.append(r)
                                    #     key_small = sheet_2.Cells(min(row_list),col).Value
                                    #     key_large = sheet_2.Cells(max(row_list), col).Value
                                    #     if 'from_key' in info['关系表']['content'][i].keys():
                                    #         log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入关系表from_key'+'\n')
                                    #         info['关系表']['content'][i]['from_key_2'] = set_tab['Id'] + '!' + key_small + ':' + key_large
                                    #     else:
                                    #         log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入关系表from_key'+'\n')
                                    #         info['关系表']['content'][i]['from_key'] = set_tab['Id'] + '!' + key_small + ':' + key_large
                                    elif '实测' in re.sub('\s*','',str(sheet_2.Cells(row, col).Value)):
                                        row_list = []
                                        for r in range(row + 1, rows_2 + 1):
                                            if re.match('%[a-zA-Z]{3,4}\d*', str(sheet_2.Cells(r, col).Value)):
                                                row_list.append(r)
                                        key_small = sheet_2.Cells(min(row_list), col).Value
                                        key_large = sheet_2.Cells(max(row_list), col).Value
                                        if 'from_key' in info['关系表']['content'][i].keys():
                                            log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入关系表from_key'+'\n')
                                            info['关系表']['content'][i]['from_key_2'] = set_tab['Id'] + '!' + key_small + ':' + key_large
                                        else:
                                            log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+ '写入关系表from_key'+'\n')
                                            info['关系表']['content'][i]['from_key'] = set_tab['Id'] + '!' + key_small + ':' + key_large

                finally:
                    work_2.Close()
                    excel.Quit()
        # print(info)
    finally:
        work.Close()
        myexcel.Quit()
    # log.write(time.strftime('%H:%M:%S', time.localtime()) +':\t'+'——————解析结束——————\n\n\n')
    log.close()
    file = open(os.path.join(log_path,time.strftime('%Y-%m-%d', time.localtime())+'key.txt'),'a',encoding='utf-8')
    file.write('\n'+time.strftime('%H:%M:%S', time.localtime()) +':\t'+str(info)+'\n')
    file.close()

    return info


# inspect_info(path, temp_path, log_path)


def evaluate_info(path):
    """获取评定表的关系信息"""
    info = {'设计值': [], '外观质量': [], '实测值': [], '质保资料':{},}
    log = open(os.path.join(log_path, time.strftime('%Y-%m-%d', time.localtime()) + 'log.txt'), 'a', encoding='utf-8')
    log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '解析：{}'.format(os.path.basename(path)) + '\n')
    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = 0
        excel.DisplayAlerts = 0
        work = excel.Workbooks.Open(path)
        sheet = work.Worksheets(1)
        rows = sheet.UsedRange.Rows.Count
        cols = sheet.UsedRange.Columns.Count
        column = None
        for row in range(1, rows+1):
            for col in range(1, cols+1):
                if re.sub('\s*','',str(sheet.Cells(row, col).Value)) == '规定值或允许偏差':
                    column = col
                    for r in range(row+1, rows+1):
                        content = {}
                        if re.match('%[a-zA-Z]{3,4}\d+',str(sheet.Cells(r, col+1).Value)):
                            for t in range(col - 1, 0, -1):
                                if sheet.Cells(r-1, t).Value:
                                    key_n = re.sub('[(（][a-zA-Z]*.?[A-Za-z]*[)）]', '',re.sub('\s*', '', str(sheet.Cells(r-1, t).Value)))
                                    log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t'+'写入设计值header'+'\n')
                                    content['header'] = key_n
                                    break
                            log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入设计值to_key' + '\n')
                            content['to_key'] = re.sub('\s*','',str(sheet.Cells(r, col+1).Value))[3:]
                            info['设计值'].append(content)
                elif re.sub('\s*','',str(sheet.Cells(row, col).Value)) == '工程质量等级评定':
                    for c in range(col+1, cols+1):
                        if re.match('%[a-zA-Z]{3,4}\d+',str(sheet.Cells(row, c).Value)):
                            log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '查找质量等级的key' + '\n')
                            info['工程质量等级评定'] = re.sub('\s*', '', str(sheet.Cells(row, c).Value))
                            break
                elif re.sub('\s*','',str(sheet.Cells(row, col).Value)) == '质量保证资料':
                    for c in range(col + 1, cols + 1):
                        if re.match('%[a-zA-Z]{3,4}\d+', str(sheet.Cells(row, c).Value)):
                            log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入质保资料to_key' + '\n')
                            info['质保资料']['to_key'] = re.sub('\s*', '', str(sheet.Cells(row, c).Value))[3:]
                            if re.search('(SG)', os.path.basename(path)):
                                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入质保资料from_key' + '\n')
                                info['质保资料']['from_key'] = 'LBZJTY0064!%CBI24'
                                break
                            elif re.search('(JL)', os.path.basename(path)):
                                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入质保资料from_key' + '\n')
                                info['质保资料']['from_key'] = 'LBZJTY0065!%CBI24'
                                break
                elif re.sub('\s*', '', str(sheet.Cells(row, col).Value)) == '外观质量':
                    count = 0
                    for c in range(col + 1, cols + 1):
                        if re.match('%[a-zA-Z]{3,4}\d+', str(sheet.Cells(row, c).Value)):
                            count += 1
                            content = {}
                            if count == 1:
                                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入外观合格to_key' + '\n')
                                content['header'] = '合格'
                                content['to_key'] = re.sub('\s*', '', str(sheet.Cells(row, c).Value))[3:]
                                info['外观质量'].append(content)
                            elif count == 2:
                                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入外观不合格to_key' + '\n')
                                content['header'] = '不合格'
                                content['to_key'] = re.sub('\s*', '', str(sheet.Cells(row, c).Value))[3:]
                                info['外观质量'].append(content)
                                break
        adr = find_adr(path, '实测值或实测偏差值')
        for row in range(adr[0]+1, rows+1):
            if re.match('%EP[a-zA-Z]{1,2}\d+', str(sheet.Cells(row, adr[1]).Value)):
                key_n = None
                start = sheet.Cells(row, adr[1]).Value.lstrip('%EP')
                end = None
                content = {}
                for col in range(column-1, 0, -1):
                    if sheet.Cells(row, col).Value:
                        key_n = re.sub('[(（][a-zA-Z]*[)）]','',re.sub('\s*','',str(sheet.Cells(row, col).Value)))
                        break
                count = 1
                for c in range(adr[1]+1, cols+1):
                    if re.match('%EP[a-zA-Z]{1,2}\d+', str(sheet.Cells(row, c).Value)):
                        count += 1
                        if count == 10:
                            end = sheet.Cells(row, c).Value.lstrip('%EP')
                            break
                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入实测值to_key' + '\n')
                content['header'] = str(key_n)
                content['to_key'] = start+'-'+end
                info['实测值'].append(content)
        f = open(os.path.join(log_path, time.strftime('%Y-%m-%d', time.localtime()) + '.txt'), 'a', encoding='utf-8')
        f.write('\n'+time.strftime('%H:%M:%S', time.localtime())+':\n\t'+os.path.basename(path)+'\n\t'+str(info)+'\n')
        f.close()
    finally:
        work.Close()
        excel.Quit()

    keyword = os.path.basename(path).split('  ')[0].lstrip('评定表')
    name = os.path.basename(os.path.dirname(path)).rstrip('评定表')
    root_path = os.path.dirname(os.path.dirname(os.path.dirname(path)))
    pre_path_ins = os.path.join(root_path, r'2 检表\{}检表'.format(name))
    pre_path_note = os.path.join(root_path, r'3 外观鉴定检查记录表\{}外观鉴定检查记录表'.format(name))
    file_list_ins = os.listdir(pre_path_ins)
    target_path_ins = ''
    for f in file_list_ins:
        if keyword in f:
            print()
            target_path_ins = os.path.join(pre_path_ins, f)
            break

    if target_path_ins:
        pass
    else:
        for f in file_list_ins:
            if keyword.split("-")[0] in f:
                target_path_ins = os.path.join(pre_path_ins, f)
                break
    target_path_note = ''
    file_list_note = os.listdir(pre_path_note)
    for f in file_list_note:
        if keyword in f:
            print(f)
            target_path_note = os.path.join(pre_path_note, f)
            break
    if target_path_note:
        pass
    else:
        for f in file_list_note:
            if keyword.split("-")[0] in f:
                target_path_note = os.path.join(pre_path_note, f)
                break


    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = 0
        excel.DisplayAlerts = 0
        work = excel.Workbooks.Open(target_path_ins)
        sheet = work.Worksheets(1)
        rows = sheet.UsedRange.Rows.Count
        cols = sheet.UsedRange.Columns.Count
        tab_set = read_path(xml_path, os.path.basename(target_path_ins))
        tar_col = None
        for row in range(1, rows+1):
            for col in range(1, cols+1):
                if re.sub('\s*','',str(sheet.Cells(row, col).Value)) == '规定值或允许偏差':
                    # print("目标列：{}".format(col))
                    tar_col = col
                    break
            if tar_col:
                break
        for row in range(1, rows+1):
            for col in range(1, cols+1):
                # if re.sub('\s*','',str(sheet.Cells(row, col).Value)) == '规定值或允许偏差':
                #     print("目标列：{}".format(col))
                #     tar_col = col
                for i in range(len(info['设计值'])):
                    if 'header' in info['设计值'][i].keys():
                        if re.sub('[a-zA-Z]{0,2}','',re.sub('[(（][a-zA-Z]*.?[a-zA-Z]*[)）]','',re.sub('\s*','',str(sheet.Cells(row, col).Value)))) == re.sub('[a-zA-Z]{0,2}','',re.sub('[(（][a-zA-Z]*.?[a-zA-Z]*[)）]','',info['设计值'][i]['header'])):
                            log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入设计值from_key' + '\n')
                            info['设计值'][i]['from_key'] = tab_set['Id']+'!'+str(sheet.Cells(row+1, tar_col+1).Value)
                for i in range(len(info['实测值'])):
                    if re.sub('[a-zA-Z]{0,2}','',re.sub('[(（][a-zA-Z]*.?[a-zA-Z]*[)）]','',re.sub('\s*', '', str(sheet.Cells(row, col).Value)))) == re.sub('[a-zA-Z]{0,2}','',re.sub('[(（][a-zA-Z]*.?[a-zA-Z]*[)）]','',info['实测值'][i]['header'])):
                        for c in range(col+1, cols+1):
                            if re.match('%[a-zA-Z]{3,4}\d+',str(sheet.Cells(row, c).Value)):
                                log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入实测值from_key' + '\n')
                                info['实测值'][i]['from_key'] = tab_set['Id']+'!'+str(sheet.Cells(row, c).Value)
    finally:
        work.Close()
        excel.Quit()
    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = 0
        excel.DisplayAlerts = 0
        work = excel.Workbooks.Open(target_path_note)
        sheet = work.Worksheets(1)
        rows = sheet.UsedRange.Rows.Count
        cols = sheet.UsedRange.Columns.Count
        tab_set = read_path(xml_path, os.path.basename(target_path_note))
        contents = []
        adr = find_adr(target_path_note, '合格判定')
        for row in range(adr[0]+1,rows+1):
            if re.match('%[a-zA-Z]{3,4}\d+',str(sheet.Cells(row, adr[1]).Value)):
                contents.append(tab_set['Id']+'!'+str(sheet.Cells(row, adr[1]).Value)+'="合格"')
        if contents != []:
            content = ','.join(contents)
            for i in range(len(info['外观质量'])):
                if info['外观质量'][i]['header'] == '合格':
                    log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入外观合格from_key' + '\n')
                    info['外观质量'][i]['from_key'] = 'IF(AND(' + content + '),"☑","☐")'
                else:
                    log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '写入外观不合格from_key' + '\n')
                    info['外观质量'][i]['from_key'] = 'IF(AND(' + content + '),"☐","☑")'
    finally:
        work.Close()
        excel.Quit()
    log.close()
    file = open(os.path.join(log_path, time.strftime('%Y-%m-%d', time.localtime()) + 'key.txt'), 'a', encoding='utf-8')
    file.write('\n' + time.strftime('%H:%M:%S', time.localtime()) + ':\n\t'+os.path.basename(path)+'\n\t' + str(info) + '\n')
    file.close()
    return info


def write_xml(path):
    """根据得到的json格式数据写到xml当中"""
    result = read_path(xml_path, os.path.basename(path))
    root_path = r'C:\Users\admin\Desktop\datarelation\result'
    new_path = os.path.join(root_path, result['XmlName'])
    dom = xml.dom.minidom.parse(new_path)
    log = open(os.path.join(log_path, time.strftime('%Y-%m-%d', time.localtime()) + 'log.txt'), 'a', encoding='utf-8')
    log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '解析：{}'.format(os.path.basename(new_path)) + '\n')
    root = dom.documentElement
    refs = dom.createElement('refs')
    root.appendChild(refs)
    title =dom.createComment('检表'+os.path.basename(path).split('  ')[0].lstrip('评定表'))
    refs.appendChild(title)
    if '评定表' in os.path.basename(path):
        info = evaluate_info(path)
        for i in range(len(info['设计值'])):
            if 'header' in info['设计值'][i].keys():
                ref = dom.createElement('ref')
                ref.setAttribute('Id', '%TC'+info['设计值'][i]['to_key'])
                ref.setAttribute('Type', 'FillInRef')
                ref.setAttribute('Range', info['设计值'][i]['to_key'])
                ref.setAttribute('FromTab', info['设计值'][i]['from_key'])
                refs.appendChild(ref)
        for i in range(len(info['实测值'])):
            if 'from_key' in info['实测值'][i].keys():
                ref = dom.createElement('ref')
                ref.setAttribute('Id', '%EP' + info['实测值'][i]['to_key'].split('-')[0])
                ref.setAttribute('Type', 'FillInRef')
                ref.setAttribute('Range', info['实测值'][i]['to_key'])
                ref.setAttribute('FromTab', info['实测值'][i]['from_key'])
                refs.appendChild(ref)
        content_1 = dom.createComment('质保资料')
        refs.appendChild(content_1)
        ref = dom.createElement('ref')
        ref.setAttribute('Id', '%TC' + info['质保资料']['to_key'])
        ref.setAttribute('Type', 'FillInRef')
        ref.setAttribute('Range', info['质保资料']['to_key'])
        ref.setAttribute('FromTab', info['质保资料']['from_key'])
        refs.appendChild(ref)
        content_2 = dom.createComment('外观质量')
        refs.appendChild(content_2)
        for i in range(len(info['外观质量'])):
            ref = dom.createElement('ref')
            ref.setAttribute('Id', '%CH' + info['外观质量'][i]['to_key'])
            ref.setAttribute('Type', 'FillInRef')
            ref.setAttribute('Range', info['外观质量'][i]['to_key'])
            ref.setAttribute('FromTab', info['外观质量'][i]['from_key'])
            refs.appendChild(ref)
    elif '检表' in os.path.basename(path):
        info = inspect_info(path, temp_path, log_path)
        for i in info.keys():
            if len(info[i]['content']) > 0:
                for j in range(len(info[i]['content'])):
                    if 'to_key' in info[i]['content'][j].keys():
                        ref = dom.createElement('ref')
                        ref.setAttribute('Id', '%EP'+info[i]['content'][j]['to_key'])
                        ref.setAttribute('Type', 'FillInRef')
                        ref.setAttribute('Range', info[i]['content'][j]['to_key'])
                        if 'from_key_2' in info[i]['content'][j].keys():
                            ref.setAttribute('FromRange', info[i]['content'][j]['from_key']+','+info[i]['content'][j]['from_key_2'])
                        else:
                            ref.setAttribute('FromRange', info[i]['content'][j]['from_key'])
                        refs.appendChild(ref)
    f = open(os.path.join(root_path, result['XmlName']),'w',encoding='utf-8')
    dom.writexml(f,addindent='\t',newl='\n',encoding='utf-8')
    f.close()
    log.write(time.strftime('%H:%M:%S', time.localtime()) + ':\t' + '解析：{}完成'.format(os.path.basename(new_path)) + '\n\n\n')
    log.close()


p_path = r'C:\Users\admin\Desktop\datarelation\result\2019年四川省公路工程施工及监理统一用表汇编\4 桥梁工程(第四册)\1 桥梁工程\1 施工单位\1 评定表\桥梁评定表\评定表8.8.1  就地浇筑拱圈评定(SG).xlsx'
# get_info(p_path,temp_path,log_path)
# print(evaluate_info(p_path))
# write_xml(path)


if __name__ == "__main__":
    # for path in paths:
    #     if os.path.isdir(path):
    #         file_list = os.listdir(path)
    #         for file in file_list:
    #             if file[-4:] == 'xlsx':
    #                 run_path = os.path.join(path, file)
    #                 write_xml(run_path)
    #     else:
    #         continue
    paths = [
        r'C:\Users\admin\Desktop\新建文件夹\datarelation\result\2019年四川省公路工程施工及监理统一用表汇编\3 路面工程(第三册)\1 路面工程\2 监理单位\1 评定表\路面评定表',
        r'C:\Users\admin\Desktop\新建文件夹\datarelation\result\2019年四川省公路工程施工及监理统一用表汇编\3 路面工程(第三册)\1 路面工程\1 施工单位\1 评定表\路面评定表',
        r'C:\Users\admin\Desktop\新建文件夹\datarelation\result\2019年四川省公路工程施工及监理统一用表汇编\4 桥梁工程(第四册)\1 桥梁工程\1 施工单位\1 评定表\桥梁评定表',
        r'C:\Users\admin\Desktop\新建文件夹\datarelation\result\2019年四川省公路工程施工及监理统一用表汇编\4 桥梁工程(第四册)\1 桥梁工程\2 监理单位\1 评定表\桥梁评定表']
    for path in paths:
        if os.path.isdir(path):
            file_list = os.listdir(path)
            for file in file_list:
                if file[-4:] == 'xlsx':
                    try:
                        excel = win32com.client.DispatchEx('Excel.Application')
                        excel.Visible = 0
                        excel.DisplayAlerts = 0
                        work = excel.Workbooks.Open(os.path.join(path,file))
                        sheet = work.Worksheets(1)
                        rows = sheet.UsedRange.Rows.Count
                        cols = sheet.UsedRange.Columns.Count
                        for row in range(1, rows+1):
                            for col in range(1, cols+1):
                                verify = 0
                                if re.sub('\s*', '', str(sheet.Cells(row, col).Value)) == '工程质量等级评定':
                                    for c in range(col + 1, cols + 1):
                                        if re.match('%[a-zA-Z]{3,4}\d+', str(sheet.Cells(row, c).Value)):
                                            pre_value = sheet.Cells(row, c).Value
                                            sheet.Cells(row, c).Value = "%CHQualified"
                                            work.Save()
                                            print('{0}  修改成功 原值为 {1}'.format(file,pre_value))
                                            verify += 1
                                            break
                                if verify:
                                    break
                                else:
                                    continue
                    finally:
                        work.Close()
                        excel.Quit()
